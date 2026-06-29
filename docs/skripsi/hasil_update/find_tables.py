import docx

def find_tables(filename):
    doc = docx.Document(filename)
    for i, t in enumerate(doc.tables):
        if len(t.rows) > 0 and len(t.columns) > 0:
            print(f"Table {i} cell 0,0: {repr(t.cell(0,0).text)}")

find_tables("001_Skripsi_Andricha Dea Mitra_Clean.docx")
