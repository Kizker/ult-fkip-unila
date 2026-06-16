import docx

doc_path = r'c:\laragon\www\ult-fkip-unila\docs\skripsi\hasil_update\001_Skripsi_Andricha Dea Mitra_Clean.docx'
doc = docx.Document(doc_path)

in_dev = False
for t in doc.tables:
    try:
        val = t.cell(0,0).text.strip()
        val2 = t.cell(0,1).text.strip()
        if 'Komentar' in val or 'Tindak Lanjut' in val or 'Aspek' in val or 'No' in val or 'Ahli' in val:
            text = " ".join([c.text.strip() for c in t.rows[0].cells])
            print(f"Header: {text}")
            text2 = " ".join([c.text.strip() for c in t.rows[1].cells]) if len(t.rows)>1 else ""
            print(f"Row 1: {text2}")
    except Exception as e:
        pass
