import docx

doc_path = r'c:\laragon\www\ult-fkip-unila\docs\skripsi\hasil_update\001_Skripsi_Andricha Dea Mitra_Clean.docx'
doc = docx.Document(doc_path)

for t in doc.tables:
    try:
        val = t.cell(0,0).text.strip()
        if 'Nama Validator' in val:
            parent = t._element.getparent()
            print("Table parent tag:", parent.tag)
            
            # Print elements around the table
            index = list(parent).index(t._element)
            print(f"Table is at index {index}")
            
            if index > 0:
                prev = list(parent)[index-1]
                print(f"Prev element: {prev.tag}")
                if prev.tag.endswith('p'):
                    p = docx.text.paragraph.Paragraph(prev, doc)
                    print(f"Prev paragraph text: {p.text}")
                    
            if index + 1 < len(list(parent)):
                nxt = list(parent)[index+1]
                print(f"Next element: {nxt.tag}")
                if nxt.tag.endswith('p'):
                    p = docx.text.paragraph.Paragraph(nxt, doc)
                    print(f"Next paragraph text: {p.text}")
    except Exception as e:
        pass
