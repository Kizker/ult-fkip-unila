import docx
from lxml import etree

doc_path = r'c:\laragon\www\ult-fkip-unila\docs\skripsi\hasil_update\001_Skripsi_Andricha Dea Mitra_Clean.docx'
doc = docx.Document(doc_path)

for t in doc.tables:
    try:
        val = t.cell(0,0).text.strip()
        if 'Nama Validator' in val:
            xml = etree.tostring(t._element, pretty_print=True).decode()
            print(xml[:500]) # Print first 500 chars to see if it looks normal
            break
    except:
        pass
